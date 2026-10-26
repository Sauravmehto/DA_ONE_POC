"""What decides whether a formula survives a change in row count.

The scanner used to judge a formula by which functions it mentioned, so
`IFERROR(SUMIFS(...), 0)` — a sum over a range with a wrapper round it — was
refused for using two functions not on a list of six. On two real client
templates that was 85% of everything a reviewer was asked to look at, which is
how a reviewer stops reading.

What actually matters is what reads the RANGES: a range summarised by an
aggregate re-extends; a range read by position does not.
"""

from __future__ import annotations

import pathlib
import re
import warnings

import pytest

from app.engine import excel_template_scanner as scanner

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

TEMPLATES = pathlib.Path(__file__).resolve().parents[2] / "test_fixtures" / "templates"


def tier(formula: str) -> str:
    return scanner.classify_formula(formula)[0]


# ── a wrapper does not change what a range means ────────────────────────────


@pytest.mark.parametrize("formula", [
    "=SUM(B3:B40)",
    "=IFERROR(SUM(B3:B40), 0)",
    "=ROUND(SUM(B3:B40), 2)",
    "=IFERROR(AVERAGEIF(LOAN_DATA!E3:E7,A17,LOAN_DATA!S3:S7),\"-\")",
    '=IFERROR(SUMIFS(Worksheet!D3:D100,Worksheet!A3:A100,">="&EDATE(A1,-6)),0)',
    '=IF(SUM(LOAN_DATA!L3:L7)*100>20,"BREACH","Compliant")',
    "=SUM(B3:B40)/COUNT(C3:C40)",
])
def test_an_aggregate_is_an_aggregate_whatever_surrounds_it(formula):
    assert tier(formula) == "tier_1", formula


@pytest.mark.parametrize("formula", [
    "=INDEX(LOAN_DATA!L3:L7,MATCH(A17,LOAN_DATA!E3:E7,0))",
    "=VLOOKUP(A1,Data!A2:C50,3,FALSE)",
    "=XLOOKUP(A1,Data!A2:A50,Data!C2:C50)",
    "=OFFSET(Data!A1,0,0,10,1)",
    "=INDIRECT(\"Data!A\"&B1)",
    "=LARGE(Data!C2:C50,3)",
    "=SUBTOTAL(9,Data!C2:C50)",
    "=SUM(INDEX(Data!A2:C50,2,3))",
])
def test_a_range_read_by_position_is_never_regenerated(formula):
    """Extending one of these changes WHICH row is read, not how many are
    summarised — the failure that shows up as a wrong number nobody queries."""
    assert tier(formula) == "tier_2", formula


@pytest.mark.parametrize("formula", [
    "=IFERROR(LOAN_DATA!L6/SUM(LOAN_DATA!L3:L7),0)",
    "=Worksheet!Q11",
    "=Worksheet!Q11*1.1",
])
def test_a_reference_pinned_to_one_record_stays_flagged(formula):
    """However clean the aggregate beside it: if the rows move, this reads a
    different record and says nothing."""
    assert tier(formula) == "tier_2", formula
    assert "pinned" in scanner.classify_formula(formula)[1] or "cross-sheet" in \
        scanner.classify_formula(formula)[1]


def test_a_bare_range_with_no_aggregate_is_flagged():
    assert tier("=B3:B40") == "tier_2"


def test_arithmetic_with_no_range_is_left_alone():
    assert tier("=B5*1.1") == "tier_3"


def test_which_ranges_an_aggregate_covers_is_reported():
    inside, outside = scanner.aggregated_ranges(
        "=IFERROR(SUMIFS(D3:D100,A3:A100,\">=\"&EDATE(A1,-6)),0)"
    )
    assert inside == ["D3:D100", "A3:A100"]
    assert outside == []

    inside, outside = scanner.aggregated_ranges("=VLOOKUP(A1,Data!A2:C50,3,FALSE)")
    assert inside == [] and outside == ["A2:C50"]


# ── against the real templates ──────────────────────────────────────────────


def _templates() -> list[pathlib.Path]:
    return sorted(TEMPLATES.glob("*.xlsx"))


@pytest.mark.skipif(not _templates(), reason="no real templates in test_fixtures/templates")
@pytest.mark.parametrize("path", _templates(), ids=lambda p: p.stem[:28])
def test_a_real_template_is_mostly_understood(path):
    """The measurement spec 0003 is judged by: how much of a real workbook the
    rules settle, and how much a person is left to read."""
    scan = scanner.scan_file(str(path))
    summary = scan["summary"]
    settled = summary["tier_1"] + summary["tier_1b_cells"]
    flagged = summary["tier_2"]

    assert settled > 0, "nothing in this workbook was understood"
    # Measured when this rule landed: 25 flagged of 12,278 formula cells in the
    # servicer report, 5 of 351 in the risk dashboard. A change that pushes
    # either back up is a regression in what a reviewer has to wade through.
    assert flagged <= 30, f"{flagged} cells flagged in {path.name}"


@pytest.mark.skipif(not _templates(), reason="no real templates in test_fixtures/templates")
@pytest.mark.parametrize("path", _templates(), ids=lambda p: p.stem[:28])
def test_every_flagged_cell_in_a_real_template_deserves_it(path):
    """Everything still flagged must be a pinned record or a positional read —
    anything else is a rule we have not written yet, not a judgement call."""
    scan = scanner.scan_file(str(path))
    for cell in scan["formula_classification"]["tier_2"]:
        reason = cell["reason"]
        assert ("pinned to specific record" in reason
                or "reads a range by position" in reason
                or "cross-sheet" in reason), f"{cell['sheet']}!{cell['cell']}: {reason}"


# ── a fill-down that skips rows, or spans columns ───────────────────


def _sheet(formulas: dict[str, str], sheet: str = "Rep") -> list[dict]:
    """The formula cells of one sheet, in the shape the detector takes."""
    cells = []
    for coord, formula in formulas.items():
        col, row = re.match(r"([A-Za-z]+)(\d+)", coord).groups()
        cells.append({"sheet": sheet, "coord": coord, "col": col,
                      "row": int(row), "formula": formula})
    return cells


def test_a_fill_down_that_prints_every_second_row_is_one_block():
    """A layout with a spacer row between records is the same fill-down with a
    gap. Judged as contiguous runs it was one flag per cell, and a reviewer
    reading forty of those reads none of them."""
    blocks = scanner.detect_row_mirror_blocks(_sheet({
        "B3": "=Src!A3", "B5": "=Src!A4", "B7": "=Src!A5", "B9": "=Src!A6",
    }))
    assert len(blocks) == 1
    assert (blocks[0]["start_row"], blocks[0]["end_row"]) == (3, 9)
    assert blocks[0]["row_step"] == 2
    assert blocks[0]["rows"] == 4


def test_a_gap_too_wide_to_be_a_layout_is_not_a_block():
    assert scanner.detect_row_mirror_blocks(_sheet({
        "B3": "=Src!A3", "B9": "=Src!A4", "B15": "=Src!A5",
    })) == []


def test_a_step_that_changes_mid_run_ends_the_run():
    """Two apart, then one apart, is two layouts meeting — not one block whose
    later rows would land in the wrong place."""
    blocks = scanner.detect_row_mirror_blocks(_sheet({
        "B3": "=Src!A3", "B5": "=Src!A4", "B7": "=Src!A5",
        "B8": "=Src!A6", "B9": "=Src!A7", "B10": "=Src!A8",
    }))
    assert {(b["start_row"], b["end_row"], b["row_step"]) for b in blocks} == {
        (3, 7, 2), (8, 10, 1)}


def test_a_stepped_run_with_jumbled_offsets_is_still_refused():
    """The step is a layout; the offsets are the meaning. Scrambled offsets are
    hand-wiring, and the whole run is refused rather than split into the part
    that happens to be clean."""
    assert scanner.detect_row_mirror_blocks(_sheet({
        "B3": "=Src!A5", "B5": "=Src!A6", "B7": "=Src!A2",
        "B9": "=Src!A3", "B11": "=Src!A4",
    })) == []


def test_a_block_advancing_more_than_one_record_a_row_is_refused():
    """Reading every second record is not a fill-down the row count can
    regenerate: extended to the data's length it would read past the end."""
    assert scanner.detect_row_mirror_blocks(_sheet({
        "B3": "=Src!A3", "B4": "=Src!A5", "B5": "=Src!A7", "B6": "=Src!A9",
    })) == []


def test_two_rows_across_adjacent_columns_are_one_region():
    """A block two rows tall is a coincidence in one column and a pattern
    across four. Judged a column at a time, all eight cells were flagged."""
    blocks = scanner.detect_row_mirror_blocks(_sheet({
        "B4": "=Src!A2", "B5": "=Src!A3",
        "C4": "=Src!B2", "C5": "=Src!B3",
        "D4": "=Src!C2*2", "D5": "=Src!C3*2",
        "E4": "=Src!D2", "E5": "=Src!D3",
    }))
    assert {b["column"] for b in blocks} == {"B", "C", "D", "E"}
    assert {b["group_id"] for b in blocks} == {"Rep!B4:E5"}
    assert all(b["group_columns"] == ["B", "C", "D", "E"] for b in blocks)


def test_two_rows_in_one_column_are_still_not_a_block():
    assert scanner.detect_row_mirror_blocks(_sheet({
        "B4": "=Src!A2", "B5": "=Src!A3"})) == []


def test_a_column_left_out_of_the_run_breaks_the_region():
    """Adjacency is the evidence. Two columns with another between them doing
    something else are two regions, and neither is wide enough."""
    assert scanner.detect_row_mirror_blocks(_sheet({
        "B4": "=Src!A2", "B5": "=Src!A3",
        "D4": "=Src!C2", "D5": "=Src!C3",
    })) == []


def test_a_tall_block_keeps_its_own_columns_grouped():
    """Grouping marks the region; it does not merge the columns, which rarely
    compute the same thing."""
    blocks = scanner.detect_row_mirror_blocks(_sheet({
        "B2": "=Src!A2", "B3": "=Src!A3", "B4": "=Src!A4",
        "C2": "=Src!B2*1.1", "C3": "=Src!B3*1.1", "C4": "=Src!B4*1.1",
    }))
    assert len(blocks) == 2
    assert {b["template_formula"] for b in blocks} == {"=Src!A2", "=Src!B2*1.1"}
    assert {b["group_id"] for b in blocks} == {"Rep!B2:C4"}


def test_a_stepped_block_is_laid_out_on_its_own_rows():
    """What the scanner records has to be what the run writes: a stepped block
    regenerated row by row would overwrite whatever the layout keeps between
    its rows."""
    import openpyxl

    from app.engine import excel_reports as er

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in (3, 5, 7):
        ws.cell(row=row, column=2).value = "=Src!A3"
    for row in (4, 6):
        ws.cell(row=row, column=2).value = "spacer"

    end = er.regenerate_block(ws, "B", 3, 7, "=Src!A3", 4, row_step=2)

    assert end == 9
    assert [ws.cell(row=r, column=2).value for r in (3, 5, 7, 9)] == [
        "=Src!A3", "=Src!A4", "=Src!A5", "=Src!A6"]
    assert [ws.cell(row=r, column=2).value for r in (4, 6)] == ["spacer", "spacer"]


def test_a_block_that_shrinks_leaves_nothing_behind():
    import openpyxl

    from app.engine import excel_reports as er

    wb = openpyxl.Workbook()
    ws = wb.active
    for k, row in enumerate((2, 3, 4, 5, 6)):
        ws.cell(row=row, column=2).value = f"=Src!A{2 + k}"

    er.regenerate_block(ws, "B", 2, 6, "=Src!A2", 2)

    assert [ws.cell(row=r, column=2).value for r in range(2, 7)] == [
        "=Src!A2", "=Src!A3", None, None, None]


# ── a template survives being renamed ───────────────────────────────


def _two_sheet_book(data_sheet: str):
    import openpyxl

    wb = openpyxl.Workbook()
    src = wb.active
    src.title = data_sheet
    src.append(["Loan", "Amount"])
    for i in range(4):
        src.append([f"L{i}", 100 + i])

    rep = wb.create_sheet("Summary")
    rep["A1"] = "Total"
    quoted = f"'{data_sheet}'" if " " in data_sheet else data_sheet
    rep["B1"] = f"=SUM({quoted}!B2:B5)"
    return wb


def test_renaming_a_sheet_does_not_make_it_a_new_template():
    """The fingerprint covered sheet names, so renaming a tab changed it in two
    places at once — the sheet list and every formula reading across to it —
    and an approved template came back unrecognised."""
    before = scanner.scan_workbook(_two_sheet_book("Raw"), "pack")
    after = scanner.scan_workbook(_two_sheet_book("Loan Tape"), "pack")

    assert before["fingerprint"] != after["fingerprint"], "the exact hash is unchanged"
    assert before["shape_fingerprint"] == after["shape_fingerprint"]


def test_a_different_shape_is_still_a_different_template():
    changed = _two_sheet_book("Raw")
    changed["Summary"]["B2"] = "=AVERAGE(Raw!B2:B5)"
    assert (scanner.scan_workbook(changed, "pack")["shape_fingerprint"]
            != scanner.scan_workbook(_two_sheet_book("Raw"), "pack")["shape_fingerprint"])


def test_a_renamed_sheet_is_told_apart_from_a_reordered_one():
    """Position is what replaces the name, so swapping two tabs is a different
    shape even though the names are the same set."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active["A1"] = "=B!A1"
    wb.create_sheet("B")["A1"] = 1

    swapped = openpyxl.Workbook()
    swapped.active.title = "B"
    swapped.active["A1"] = 1
    swapped.create_sheet("A")["A1"] = "=B!A1"

    assert (scanner.scan_workbook(wb, "t")["shape_fingerprint"]
            != scanner.scan_workbook(swapped, "t")["shape_fingerprint"])


# ── which sheet holds the records ───────────────────────────────────────────


def test_a_templates_data_sheet_is_recognised_with_few_sample_rows():
    """A template's data sheet is REFILLED at run time, so it ships with
    whatever sample rows somebody left in it — the servicer report has two.
    Requiring three made the Report Builder call it a parameter sheet, and a
    report cannot bind a parameter sheet, so the template could not be used
    at all."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    headers = [f"FIELD_{i}" for i in range(12)]
    ws.append(headers)
    ws.append(["v"] * 12)
    ws.append(["v"] * 12)

    scan = scanner.scan_workbook(wb, "sample")
    assert scan["tab_roles"]["Raw Data"] == "data"


def test_a_label_and_value_sheet_is_still_a_parameter_sheet():
    """The rule above must not turn a fund-parameters sheet into a data one:
    it is narrow, and its value column mixes types."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASSUMPTIONS"
    for label, value in [("Fund", "Bridge II"), ("Vintage", 2024), ("Rate", 0.085),
                         ("Currency", "USD")]:
        ws.append([label, value])

    scan = scanner.scan_workbook(wb, "sample")
    assert scan["tab_roles"]["ASSUMPTIONS"] == "parameter"


@pytest.mark.skipif(not _templates(), reason="no real templates in test_fixtures/templates")
def test_the_servicer_report_knows_where_its_data_lives():
    path = next((p for p in _templates() if "Meridian" in p.name), None)
    if path is None:
        pytest.skip("the servicer report is not among the templates")
    roles = scanner.scan_file(str(path))["tab_roles"]
    assert roles.get("Raw Data") == "data", roles
